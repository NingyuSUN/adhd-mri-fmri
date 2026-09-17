"""Build a results-only ADHD deck from reviewed content and frozen figure/table inputs.
Usage: python build.py --repo REPO --output-dir OUTPUT
Requires python-pptx and Pillow. No model fitting or source deck modification.
"""
import argparse,json,math,hashlib,shutil,csv,zipfile
from pathlib import Path
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION,XL_LEGEND_POSITION,XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from PIL import Image

FONT='Microsoft YaHei';NAVY='132C40';TEAL='087F8C';GRAY='536574';BG='F5F7FA';ORANGE='9B5510'
def rgb(s):return RGBColor.from_string(s)
def box(sl,x,y,w,h,fill):
 s=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));s.fill.solid();s.fill.fore_color.rgb=rgb(fill);s.line.fill.background();return s

def tx(sl,x,y,w,h,text,size=18,color=NAVY,bold=False):
 s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(6)
 return s

def width_units(s):return sum(1 if ord(c)>255 else .53 for c in s)
def table(sl,rows,x=.62,y=1.94,w=12.1,h=4.4):
 nr,nc=len(rows),len(rows[0]);size=17 if nr<=7 else 15.5
 if nc>=5:size=15.5
 weights=[]
 for j in range(nc):
  lens=[width_units(str(r[j])) for r in rows]
  weights.append(min(25,max(9,max(lens))))
 if nc==4 and 'AUC' in ''.join(rows[0]):weights=[25,16,17,17]
 ws=[w*v/sum(weights) for v in weights]
 # Avoid narrow prose columns; rows grow with wrapped content.
 units=[]
 for i,r in enumerate(rows):
  lines=max(math.ceil(width_units(str(v))/(max(.3,ws[j]-.18)*72/size)) for j,v in enumerate(r))
  units.append(max(1,lines)*1.1+.3)
 if sum(units)*size/72>h:size=max(13.5,size-1)
 shape=sl.shapes.add_table(nr,nc,Inches(x),Inches(y),Inches(w),Inches(h));t=shape.table
 for j,cw in enumerate(ws):t.columns[j].width=Inches(cw)
 for i in range(nr):
  t.rows[i].height=Inches(h*units[i]/sum(units))
  for j in range(nc):
   c=t.cell(i,j);c.text=str(rows[i][j]);c.fill.solid();c.fill.fore_color.rgb=rgb(NAVY if i==0 else ('FFFFFF' if i%2 else 'EAF0F4'));c.margin_left=Inches(.1);c.margin_right=Inches(.08);c.margin_top=c.margin_bottom=Inches(.035);c.vertical_anchor=MSO_ANCHOR.MIDDLE
   for p in c.text_frame.paragraphs:
    p.font.name=FONT;p.font.size=Pt(size);p.font.bold=i==0;p.font.color.rgb=rgb('FFFFFF' if i==0 else NAVY);p.space_after=Pt(0)
 return shape

def picture_fit(sl,p,x,y,w,h):
 iw,ih=Image.open(p).size;r=min(w/iw,h/ih);ww,hh=iw*r,ih*r
 sl.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

def chart(sl,d):
 cd=CategoryChartData();cd.categories=d['categories']
 for s in d['series']:cd.add_series(s['name'],s['values'])
 shape=sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(.65), Inches(1.95), Inches(12), Inches(4.35),cd);c=shape.chart
 c.has_legend=len(d['series'])>1
 if c.has_legend:c.legend.position=XL_LEGEND_POSITION.BOTTOM;c.legend.font.name=FONT;c.legend.font.size=Pt(14);c.legend.include_in_layout=False
 c.value_axis.minimum_scale=0;c.value_axis.maximum_scale=1;c.value_axis.major_unit=.2;c.value_axis.tick_labels.number_format='0.0';c.value_axis.tick_labels.font.size=Pt(13);c.value_axis.tick_labels.font.name=FONT
 c.value_axis.has_major_gridlines=True;c.category_axis.tick_labels.font.name=FONT;c.category_axis.tick_labels.font.size=Pt(15);c.category_axis.reverse_order=True
 c.category_axis.major_tick_mark=XL_TICK_MARK.NONE
 c.plots[0].has_data_labels=True;c.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;c.plots[0].data_labels.number_format='0.000';c.plots[0].data_labels.font.name=FONT;c.plots[0].data_labels.font.size=Pt(14)
 for i,s in enumerate(c.series):s.format.fill.solid();s.format.fill.fore_color.rgb=rgb(TEAL if i==0 else 'DA884B');s.format.line.fill.background()
 # Normalize python-pptx legacy signed axis IDs to OOXML unsigned IDs.
 for el in c._chartSpace.iter():
  if el.tag.rsplit('}',1)[-1] in ('axId','crossAx') and el.get('val'):
   el.set('val',str(int(el.get('val')) % (2**32)))
 return shape

def build(repo,out):
 data=json.loads((Path(__file__).parent/'content.json').read_text(encoding='utf8'));out.mkdir(parents=True,exist_ok=True)
 src=repo/'results/tables_figures_20260914';pr=Presentation();pr.slide_width=Inches(13.333333);pr.slide_height=Inches(7.5)
 for d in data['slides']:
  sl=pr.slides.add_slide(pr.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG)
  box(sl,0,0,13.333,.1,TEAL);tx(sl,.62,.23,12,.24,d['group']+'  /  ADHD-200',11,TEAL,True)
  tx(sl,.62,.63,12.1,.56,d['title'],27,NAVY,True);tx(sl,.62,1.3,12.1,.38,d['sub'],14,GRAY)
  if d.get('warning'):box(sl,.62,1.7,12.1,.25,'FFF0D9');tx(sl,.73,1.72,11.9,.2,d['warning'],11,ORANGE,True)
  if d.get('image'):
   if d.get('rows_side'):
    table(sl,d['rows_side'],.62,2.35,6.4,3.55);picture_fit(sl,src/'figures'/d['image'],7.25,1.98,5.5,4.42)
   else:picture_fit(sl,src/'figures'/d['image'],.65,1.98,12.05,4.42)
  elif d.get('chart'):chart(sl,d['chart'])
  elif d.get('rows'):table(sl,d['rows'])
  else:
   bullets=d.get('bullets') or []
   gap=min(1.18,4.25/max(1,len(bullets)))
   for i,b in enumerate(bullets):
    yy=2.05+i*gap;box(sl,.65,yy+.07,.08,.26,TEAL);tx(sl,.9,yy,11.55,gap-.1,b,23 if len(bullets)<=3 else 21)
  box(sl,.62,6.58,12.1,.46,'E1EFF0');tx(sl,.78,6.665,11.8,.28,d['q'],14,TEAL,True)
  origin=d.get('origin');label=('原 fMRI' if origin and origin['deck'].startswith('ADHD_fMRI') else '原多模态') if origin else '综合整理'
  source=(f'{label} 第{origin["slide"]}页' if origin else label)
  if d.get('source','').startswith(('tables/','figures/')):source+=' | '+d['source']
  tx(sl,.62,7.16,11.2,.19,'来源：'+source,9,GRAY);tx(sl,12,7.12,.7,.25,f'{d["slide"]:02d} / {len(data["slides"])}',11,GRAY)
  sl.notes_slide.notes_text_frame.text=d['notes']
 pr.core_properties.title='ADHD 结果与论文主线：详细备注版';pr.core_properties.subject='Results extracted from original fMRI and multimodal presentations; evidence boundaries and paper framing';pr.core_properties.author='ADHD project';pr.core_properties.keywords='ADHD; MRI; fMRI; QC; LOSO; incremental prediction'
 path=out/'ADHD_结果与论文主线_详细备注版.pptx';pr.save(path)
 # Companion notes and source-to-output map, kept readable independently of PowerPoint.
 lines=['# ADHD 结果与论文主线：逐页详细讲稿','',f'主线：第1–{data["metadata"]["main_slides"]}页；补充结果：其后各页。','']
 for d in data['slides']:lines+=['## 第'+str(d['slide'])+'页｜'+d['title'],'',d['notes'],'']
 (out/'逐页详细讲稿.md').write_text('\n'.join(lines),encoding='utf8')
 with (out/'原课件到新PPT页码.csv').open('w',newline='',encoding='utf-8-sig') as f:
  w=csv.writer(f);w.writerow(['新页码','标题','原课件','原页码','数据来源'])
  for d in data['slides']:
   o=d.get('origin') or {};w.writerow([d['slide'],d['title'],o.get('deck','新增综合页'),o.get('slide',''),d['source']])
 for sub in ['tables','figures']:
  (out/sub).mkdir(exist_ok=True)
  for p in (src/sub).iterdir():
   if p.is_file():shutil.copy2(p,out/sub/p.name)
 shutil.copy2(Path(__file__).parent/'content.json',out/'逐页内容与来源.json')
 readme='''# ADHD 结果专册（2026-09-17）\n\n打开 PPTX，在 PowerPoint 的“备注”或“演讲者视图”阅读每页详解。逐页详细讲稿.md 提供同样说明。\n\n- 主线 26 页，补充结果 27 页，共53页。\n- 原始来源：同一ADHD-ppt文件夹的ADHD_fMRI.pptx与2026-09-14多模态图表版；原文件未修改。\n- 14份完整CSV表、9幅PNG科研图及已有PDF随附；长表的PPT页为原课件摘要。\n- 原始LOSO区间保留历史值，并标注同分算法实际影响待核验；没有重算或声称修复。\n- 本轮仅提取、重排、解释既有证据；没有新训练、原始MRI QC或外部验证。\n- GitHub当前main：3f39504；本地结果依据：ed50a60。冻结图表目录经Git比较无差异。本地文献/AI审阅增补不冒充GitHub已发布成果。\n- 旧fMRI柱图由原课件缓存值重建为可编辑图表，将AUC×100统一除以100。\n- 原始PPT来源SHA-256、每页来源见逐页内容与来源.json；新旧页码见CSV。\n\n推荐论文角度：多站点ADHD影像融合的增量预测及其对质量控制、内部选择与评估方式的敏感性。它是建议，不是已确认的新颖性或发表保证。\n'''
 (out/'README_先读我.md').write_text(readme,encoding='utf8')
 # Read back actual notes and OPC package; check that every substantive slide has all notes.
 reread=Presentation(path);checks=[]
 for i,s in enumerate(reread.slides):
  expected=data['slides'][i];actual=s.notes_slide.notes_text_frame.text
  assert actual==expected['notes'],f'notes mismatch {i+1}'
  assert len(actual)>250,f'notes short {i+1}'
  checks.append({'slide':i+1,'notes_chars':len(actual),'tables':sum(sh.has_table for sh in s.shapes),'charts':sum(sh.has_chart for sh in s.shapes)})
 assert len(reread.slides)==53
 assert len(list((out/'tables').glob('*.csv')))==14
 assert len(list((out/'figures').glob('*.png')))==9
 with zipfile.ZipFile(path) as z:assert z.testzip() is None
 audit={'slide_count':len(reread.slides),'main_slides':26,'notes_match':True,'notes_chars':sum(x['notes_chars'] for x in checks),'native_tables':sum(x['tables'] for x in checks),'native_charts':sum(x['charts'] for x in checks),'source_metadata':data['metadata'],'per_slide':checks,'pptx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'visual_review':'pending rendering'}
 (out/'validation.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps({'output':str(path),'slides':len(reread.slides),'notes_chars':audit['notes_chars'],'tables':audit['native_tables'],'charts':audit['native_charts']},ensure_ascii=True))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--repo',type=Path,required=True);ap.add_argument('--output-dir',type=Path,required=True);a=ap.parse_args();build(a.repo,a.output_dir)
